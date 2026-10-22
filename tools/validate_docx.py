"""Valida data/funds.json contra el Word de origen. Sale con error si algo no coincide.

Uso:  python tools/validate_docx.py ruta/al/documento.docx
"""
import os, re, sys, json, unicodedata

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.environ.get('CFO_SRC') or (sys.argv[1] if len(sys.argv) > 1 else '')
if not SRC:
    sys.exit(__doc__)

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

d = docx.Document(SRC)

def body_items(doc):
    for ch in doc.element.body.iterchildren():
        if ch.tag.endswith('}p'):
            yield Paragraph(ch, doc)
        elif ch.tag.endswith('}tbl'):
            yield Table(ch, doc)

seq = list(body_items(d))

def flat(t):
    out = []
    for r in t.rows:
        for c in r.cells:
            out += [p.text.strip() for p in c.paragraphs if p.text.strip()]
            for nt in c.tables:
                out += flat(nt)
    return out

def is_head(n):
    it = seq[n]
    return (isinstance(it, Paragraph) and it.style.name == 'Heading 1'
            and any(isinstance(seq[m], Paragraph) and seq[m].text.strip().startswith('ISIN ')
                    for m in range(n + 1, min(n + 4, len(seq)))))

heads = [n for n in range(len(seq)) if is_head(n)]
blocks = {}
for k, n in enumerate(heads):
    end = heads[k + 1] - 1 if k + 1 < len(heads) else len(seq)
    txt = []
    for it in seq[n:end]:
        txt += [it.text.strip()] if isinstance(it, Paragraph) else flat(it)
    isin = next(t.split()[1] for t in txt if t.startswith('ISIN '))
    blocks[isin] = [t for t in txt if t]

maptext = []
for n in range(heads[0]):
    if isinstance(seq[n], Table):
        maptext += flat(seq[n])

data = json.load(open(os.path.join(ROOT, 'data', 'funds.json')))
F = data['funds']
err = []

def norm(s):
    s = unicodedata.normalize('NFC', str(s)).replace(' ', ' ')
    return re.sub(r'\s+', ' ', s).strip()

def seen(val, pool):
    """El valor aparece en el bloque, con coma o punto decimal."""
    v = norm(val)
    alt = {v, v.replace(',', '.'), v.replace('.', ','), v.rstrip('%'), v + '%'}
    pool = {norm(x) for x in pool}
    return bool(alt & pool) or any(any(a and a in x for a in alt) for x in pool)

if len(F) != len(blocks):
    err.append(f'funds.json tiene {len(F)} fondos y el documento {len(blocks)}')
if len({f['isin'] for f in F}) != len(F):
    err.append('hay ISIN duplicados')

for f in F:
    b = blocks.get(f['isin'])
    if not b:
        err.append(f'{f["isin"]}: no está en el documento')
        continue
    pool = b + maptext
    if f['objective'][1:] not in ' '.join(b):
        err.append(f'{f["isin"]}: el objetivo no coincide literalmente')
    if not seen(f['aum'], b):
        err.append(f'{f["isin"]}: AUM {f["aum"]} no figura en la ficha')
    for k, r in f['returns'].items():
        if r and not seen(f'{r["v"]:.{r["d"]}f}%'.replace('.', ','), b):
            err.append(f'{f["isin"]}: rendimiento {k}={r} no figura')
    for key in ('ter', 'trailer'):
        r = f[key]
        if r and not seen(f'{r["v"]:.{r["d"]}f}%'.replace('.', ','), pool):
            err.append(f'{f["isin"]}: {key} {r} no figura')
    for s in f['sections']:
        for it in s['items']:
            if not seen(it['label'][0].lower() + it['label'][1:], b) and not seen(it['label'], b):
                err.append(f'{f["isin"]}: etiqueta "{it["label"]}" no figura')
            if it['value'] and not seen(it['value'], b):
                err.append(f'{f["isin"]}: valor "{it["value"]}" de "{it["label"]}" no figura')
    for n in f.get('notes', []):
        if n not in b:
            err.append(f'{f["isin"]}: nota no literal')
    plats = [p for p, ok in f['platforms'].items() if ok]
    linea = next((x for x in b if x.startswith(('BCV', 'BCI', 'Stonex'))), '')
    if len(plats) != len([x for x in linea.split(',') if x.strip()]):
        err.append(f'{f["isin"]}: plataformas no coinciden ({linea})')
    if f['asset'] == 'rf' and not f['balanz'] and f['ytm'] is None:
        err.append(f'{f["isin"]}: fondo de renta fija sin YTM')

html_path = os.path.join(ROOT, 'index.html')
if os.path.exists(html_path):
    html = open(html_path).read()
    if 'Nota sobre los datos' in html:
        err.append('index.html contiene "Nota sobre los datos de la fuente"')

print('ERRORES:', err if err else 'ninguno')
print(f'Fondos: {len(F)} | propios: {sum(f["balanz"] for f in F)} | '
      f'categorías: {len(data["categories"])}')
if err:
    sys.exit(1)
